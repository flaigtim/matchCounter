from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font
from openpyxl.styles import PatternFill


# Farben fuer Ehrungen
HONOR_REACHED_FONT = Font(color="9C0006", bold=True)  # Rot fuer erreichte Ehrung
HONOR_POTENTIAL_FONT = Font(color="FF8C00", bold=True)  # Orange fuer potenzielle Ehrung


# =========================================
# KONFIGURATION (hier jaehrlich anpassen)
# =========================================
YEARLY_HERREN_FILE = "player_stats/yearly_stats/2025_2026/herren_stats_2025_2026.xlsx"
OVERALL_HERREN_FILE = "player_stats/_herren_stats_overall.xlsx"
YEARLY_DAMEN_FILE = "player_stats/yearly_stats/2025_2026/damen_stats_2025_2026.xlsx"
OVERALL_DAMEN_FILE = "player_stats/_damen_stats_overall.xlsx"
NEW_SHEET    = "25-26"   # Name des neuen Tabellenblatts
PREV_SHEET   = "24-25"   # Vorsaison-Blatt, das kopiert wird


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Fuehrt den Saison-Merge fuer Herren und Damen durch."
    )
    parser.add_argument("--yearly-herren", default=YEARLY_HERREN_FILE, help=f"Yearly-Datei Herren (Standard: {YEARLY_HERREN_FILE})")
    parser.add_argument("--overall-herren", default=OVERALL_HERREN_FILE, help=f"Overall-Datei Herren (Standard: {OVERALL_HERREN_FILE})")
    parser.add_argument("--yearly-damen", default=YEARLY_DAMEN_FILE, help=f"Yearly-Datei Damen (Standard: {YEARLY_DAMEN_FILE})")
    parser.add_argument("--overall-damen", default=OVERALL_DAMEN_FILE, help=f"Overall-Datei Damen (Standard: {OVERALL_DAMEN_FILE})")
    parser.add_argument("--new-sheet", default=NEW_SHEET, help=f"Neues Blatt (Standard: {NEW_SHEET})")
    parser.add_argument("--prev-sheet", default=PREV_SHEET, help=f"Vorsaison-Blatt (Standard: {PREV_SHEET})")
    return parser.parse_args()


def normalize_name(value: str) -> str:
    return " ".join(value.strip().casefold().split())


def to_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def read_yearly_players(yearly_path: Path) -> list[tuple[str, int, int, int]]:
    if not yearly_path.exists():
        raise FileNotFoundError(f"Yearly-Datei nicht gefunden: {yearly_path}")

    yearly_wb = openpyxl.load_workbook(yearly_path, data_only=True)
    sheet = yearly_wb["Zusammenfassung"] if "Zusammenfassung" in yearly_wb.sheetnames else yearly_wb.active

    # Bevorzugt explizite Spalten aus der Kopfzeile.
    name_col = None
    punkt_col = None
    pokal_col = None
    test_col = None
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=col).value
        if not value:
            continue
        header = str(value).strip().lower()
        if header in {"spieler", "name", "spielername"}:
            name_col = col
        if header in {"punktspiel", "punktspiele", "puntkspiel"}:
            punkt_col = col
        if header in {"pokalspiel", "pokalspiele"}:
            pokal_col = col
        if header in {"testspiel", "testspiele", "freundschaftsspiel", "freundschaftsspiele"}:
            test_col = col

    if name_col is None:
        name_col = 1
    if punkt_col is None:
        raise ValueError("In der yearly-Datei wurde keine Spalte 'Punktspiel' gefunden.")
    if pokal_col is None:
        raise ValueError("In der yearly-Datei wurde keine Spalte 'Pokalspiel' gefunden.")
    if test_col is None:
        raise ValueError("In der yearly-Datei wurde keine Spalte 'Testspiel' gefunden.")

    players: list[tuple[str, int, int, int]] = []
    seen: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        name_value = sheet.cell(row=row, column=name_col).value
        if not name_value:
            continue
        name = str(name_value).strip()
        if not name:
            continue
        key = name.casefold()
        if key in seen:
            continue
        seen.add(key)
        punkt = to_int(sheet.cell(row=row, column=punkt_col).value)
        pokal = to_int(sheet.cell(row=row, column=pokal_col).value)
        test = to_int(sheet.cell(row=row, column=test_col).value)
        players.append((name, punkt, pokal, test))

    return players


def read_manual_players(manual_path: Path) -> dict[str, tuple[str, int]]:
    """Liest optional manuell erfasste Spiele aus tournaments.xlsx (A: Name, B: Spiele)."""
    if not manual_path.exists():
        return {}

    manual_wb = openpyxl.load_workbook(manual_path, data_only=True)
    sheet = manual_wb.active

    players: dict[str, tuple[str, int]] = {}
    for row in range(2, sheet.max_row + 1):  # Zeile 1 = Kopfzeile
        name_value = sheet.cell(row=row, column=1).value
        if not name_value:
            continue
        name = str(name_value).strip()
        if not name:
            continue

        games = to_int(sheet.cell(row=row, column=2).value)
        key = normalize_name(name)
        if key in players:
            existing_name, existing_games = players[key]
            players[key] = (existing_name, existing_games + games)
        else:
            players[key] = (name, games)

    return players


def read_prev_name_rows(prev_ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, int]:
    first_data_row = 8
    name_to_row: dict[str, int] = {}

    for row in range(first_data_row, prev_ws.max_row + 1):
        value = prev_ws.cell(row=row, column=1).value
        if not value:
            continue
        key = normalize_name(str(value))
        if not key or key in name_to_row:
            continue
        name_to_row[key] = row

    return name_to_row


def read_thresholds(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[int]:
    """Liest Schwellwerte aus J3:R3."""
    thresholds: list[int] = []
    for col in range(10, 19):  # J bis R
        value = ws.cell(row=3, column=col).value
        if value is None:
            break
        thresholds.append(to_int(value))
    return thresholds


def merge_one(yearly_path: Path, overall_path: Path, new_sheet: str, prev_sheet: str, label: str) -> None:

    if not overall_path.exists():
        raise FileNotFoundError(f"Overall-Datei nicht gefunden: {overall_path}")

    overall_wb = openpyxl.load_workbook(overall_path)

    if prev_sheet not in overall_wb.sheetnames:
        raise ValueError(
            f"Vorsaison-Blatt '{prev_sheet}' nicht gefunden. "
            f"Vorhanden: {', '.join(overall_wb.sheetnames)}"
        )

    if new_sheet in overall_wb.sheetnames:
        del overall_wb[new_sheet]

    src = overall_wb[prev_sheet]
    prev_name_rows = read_prev_name_rows(src)
    ws = overall_wb.copy_worksheet(src)
    ws.title = new_sheet
    ws.freeze_panes = "B8"

    yearly_players = read_yearly_players(yearly_path)
    manual_path = yearly_path.parent / "tournaments.xlsx"
    manual_players = read_manual_players(manual_path)

    merged_players: list[tuple[str, int, int, int, int]] = []
    used_manual_keys: set[str] = set()

    for name, punktspiel, pokalspiel, testspiel in yearly_players:
        key = normalize_name(name)
        manual_games = 0
        if key in manual_players:
            manual_games = manual_players[key][1]
            used_manual_keys.add(key)
        merged_players.append((name, punktspiel, pokalspiel, testspiel, manual_games))

    for key, (name, manual_games) in manual_players.items():
        if key in used_manual_keys:
            continue
        # Spieler nur aus tournaments.xlsx ebenfalls im Gesamtblatt aufnehmen.
        merged_players.append((name, 0, 0, 0, manual_games))

    # Ausgabe im neuen Saisonblatt alphabetisch nach Name (Spalte A) sortieren.
    merged_players.sort(key=lambda player: normalize_name(player[0]))

    # B1:D1 mit Saisontext aus dem Blattnamen fuellen (z. B. 25-26 -> Saison 2025/26).
    try:
        season_xx, season_yy = new_sheet.split("-", 1)
    except ValueError as exc:
        raise ValueError(
            f"new-sheet muss das Format 'xx-yy' haben, erhalten: {new_sheet}"
        ) from exc
    ws["B1"] = f"Saison 20{season_xx}/{season_yy}"
    ws["B4"] = f"Stand {datetime.now().strftime('%d.%m.%Y')}"

    # Bereiche A8:F<letzte Zeile>, G8:I<letzte Zeile> leeren.
    first_data_row = 8
    last_row = max(ws.max_row, first_data_row + len(merged_players) - 1)
    for row in range(first_data_row, last_row + 1):
        for col in range(1, 10):  # A bis I
            ws.cell(row=row, column=col).value = None

    # Schwellwerte lesen.
    thresholds = read_thresholds(ws)

    # Namen aus der yearly-Datei in Spalte A schreiben.
    missing_prev_rows: list[int] = []
    g_values: dict[int, int] = {}
    
    for idx, (name, punktspiel, pokalspiel, testspiel, manual_games) in enumerate(merged_players):
        target_row = first_data_row + idx
        ws.cell(row=target_row, column=1).value = name
        ws.cell(row=target_row, column=3).value = punktspiel
        ws.cell(row=target_row, column=4).value = pokalspiel
        ws.cell(row=target_row, column=5).value = testspiel
        ws.cell(row=target_row, column=6).value = manual_games

        # Spalte B: Vorsaison-Spiele
        prev_row = prev_name_rows.get(normalize_name(name))
        b_value = 0
        if prev_row is not None:
            b_value = to_int(src.cell(row=prev_row, column=8).value)
            ws.cell(row=target_row, column=2).value = f"='{prev_sheet}'!H{prev_row}"
        else:
            b_cell = ws.cell(row=target_row, column=2)
            b_cell.value = "=0"
            missing_prev_rows.append(target_row)

        # Spalten G und H: Als Formeln
        ws.cell(row=target_row, column=7).value = f"=C{target_row}+D{target_row}+E{target_row}+F{target_row}"
        ws.cell(row=target_row, column=8).value = f"=G{target_row}+B{target_row}"

        # Berechne G/H lokal fuer Ehrungslogik und Formatierung
        g_value = punktspiel + pokalspiel + testspiel + manual_games
        h_value = g_value + b_value
        g_values[target_row] = g_value

        # Spalte I: Ehrung mit Formel
        ws.cell(row=target_row, column=9).value = f"=IFERROR(LOOKUP(2,1/((B{target_row}<$J$3:$R$3)*(H{target_row}>=$J$3:$R$3)),$J$3:$R$3),\"\")"

    # Datenbereich ab Zeile 8 abwechselnd einfärben.
    fill_a = PatternFill(fill_type="solid", fgColor="FFFFFF")
    fill_b = PatternFill(fill_type="solid", fgColor="EDEDED")
    missing_font = Font(color="9C0006", bold=True)
    for row in range(first_data_row, last_row + 1):
        fill = fill_a if (row - first_data_row) % 2 == 0 else fill_b
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).fill = fill

    # Fehlende Vorjahreswerte in Spalte B sichtbar markieren (rot).
    for row in missing_prev_rows:
        ws.cell(row=row, column=2).font = HONOR_REACHED_FONT

    # Alle Werte in Spalte I rot und fett markieren.
    for row in range(first_data_row, last_row + 1):
        ws.cell(row=row, column=9).font = HONOR_REACHED_FONT

    # Spalte G >= 50 orange und fett markieren.
    for row, g_val in g_values.items():
        if g_val >= 50:
            ws.cell(row=row, column=7).font = HONOR_POTENTIAL_FONT

    overall_wb.save(overall_path)
    print(f"[{label}] Blatt '{prev_sheet}' kopiert nach '{new_sheet}'")
    print(f"[{label}] Namen aus yearly uebernommen: {len(yearly_players)}")
    if manual_path.exists():
        print(f"[{label}] tournaments.xlsx gefunden: {manual_path}")
        print(f"[{label}] Zusatzeintraege aus tournaments.xlsx: {len(manual_players)}")
    else:
        print(f"[{label}] Keine tournaments.xlsx gefunden (optional): {manual_path}")
    print(f"[{label}] Namen gesamt im neuen Blatt: {len(merged_players)}")
    print(f"[{label}] Gespeichert: {overall_path}")


def main() -> None:
    args = parse_args()
    new_sheet = args.new_sheet
    prev_sheet = args.prev_sheet

    merge_one(
        yearly_path=Path(args.yearly_herren),
        overall_path=Path(args.overall_herren),
        new_sheet=new_sheet,
        prev_sheet=prev_sheet,
        label="Herren",
    )

    merge_one(
        yearly_path=Path(args.yearly_damen),
        overall_path=Path(args.overall_damen),
        new_sheet=new_sheet,
        prev_sheet=prev_sheet,
        label="Damen",
    )


if __name__ == "__main__":
    main()
