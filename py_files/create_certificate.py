from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl


OVERALL_FILE = "player_stats/_stats_overall.xlsx"
DEFAULT_SHEET = "25-26"
DEFAULT_TEMPLATE = "template_certificate.docx"
DEFAULT_OUTPUT_DIR = "player_stats/certificates/" + DEFAULT_SHEET
DEFAULT_CERTIFICATE_DATE = "06.06.2026"
FIRST_DATA_ROW = 8                  # Zeile in Excel-Datei, ab der die Spielerdaten beginnen (z. B. 8)


def to_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def split_name(full_name: str) -> tuple[str, str]:
    """Spaltet Namen im Format 'Nachname, Vorname' in (Vorname, Nachname)."""
    text = full_name.strip()
    if not text:
        return "", ""
    if "," in text:
        # Format: "Nachname, Vorname"
        parts = text.split(",", 1)
        last_name = parts[0].strip()
        first_name = parts[1].strip() if len(parts) > 1 else ""
        return first_name, last_name
    # Fallback für "Vorname Nachname" Format
    parts = text.split()
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def make_safe_filename(name: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._ -]", "_", name.strip())
    safe = re.sub(r"\s+", "_", safe)
    return safe or "certificate"


def normalize_name_part(value: str) -> str:
    text = value.strip()
    translit = {
        "ä": "ae",
        "ö": "oe",
        "ü": "ue",
        "Ä": "Ae",
        "Ö": "Oe",
        "Ü": "Ue",
        "ß": "ss",
    }
    for old, new in translit.items():
        text = text.replace(old, new)
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[^A-Za-z0-9]", "", text)
    return text


def read_thresholds(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[int]:
    thresholds: list[int] = []
    for col in range(10, 19):  # J bis R
        value = ws.cell(row=3, column=col).value
        if value is None:
            continue
        thresholds.append(to_int(value))
    return sorted([t for t in thresholds if t > 0])


def reached_honors(b_value: int, h_value: int, thresholds: list[int]) -> list[int]:
    reached: list[int] = []
    for threshold in thresholds:
        if b_value < threshold <= h_value:
            reached.append(threshold)
    return reached


def parse_h_reference(formula: str) -> tuple[str, int] | None:
    match = re.match(r"^='([^']+)'!H(\d+)$", formula)
    if not match:
        match = re.match(r"^=([^!]+)!H(\d+)$", formula)
    if not match:
        return None
    return match.group(1), int(match.group(2))


def resolve_b_value(
    wb: openpyxl.Workbook,
    ws: openpyxl.worksheet.worksheet.Worksheet,
    row: int,
    cache: dict[tuple[str, int], int],
    stack: set[tuple[str, int]],
) -> int:
    value = ws.cell(row=row, column=2).value
    if not isinstance(value, str):
        return to_int(value)

    formula = value.strip()
    if formula == "=0":
        return 0

    reference = parse_h_reference(formula)
    if reference is None:
        return to_int(value)

    source_sheet, source_row = reference
    if source_sheet not in wb.sheetnames:
        return 0

    return compute_h_value(wb, wb[source_sheet], source_row, cache, stack)


def compute_h_value(
    wb: openpyxl.Workbook,
    ws: openpyxl.worksheet.worksheet.Worksheet,
    row: int,
    cache: dict[tuple[str, int], int],
    stack: set[tuple[str, int]],
) -> int:
    key = (ws.title, row)
    if key in cache:
        return cache[key]
    if key in stack:
        return 0

    stack.add(key)
    b_value = resolve_b_value(wb, ws, row, cache, stack)
    g_value = (
        to_int(ws.cell(row=row, column=3).value)
        + to_int(ws.cell(row=row, column=4).value)
        + to_int(ws.cell(row=row, column=5).value)
        + to_int(ws.cell(row=row, column=6).value)
    )
    h_value = b_value + g_value
    cache[key] = h_value
    stack.remove(key)
    return h_value


def iter_honored_players(
    overall_file: Path,
    sheet_name: str,
) -> list[tuple[str, int]]:
    if not overall_file.exists():
        raise FileNotFoundError(f"Overall-Datei nicht gefunden: {overall_file}")

    wb = openpyxl.load_workbook(overall_file, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Tabellenblatt '{sheet_name}' nicht gefunden. "
            f"Vorhanden: {', '.join(wb.sheetnames)}"
        )

    ws = wb[sheet_name]
    thresholds = read_thresholds(ws)
    if not thresholds:
        raise ValueError("Keine Schwellwerte in J3:R3 gefunden.")

    honored: list[tuple[str, int]] = []
    h_cache: dict[tuple[str, int], int] = {}
    row = FIRST_DATA_ROW
    while row <= ws.max_row:
        name_value = ws.cell(row=row, column=1).value
        if not name_value:
            row += 1
            continue

        full_name = str(name_value).strip()
        if not full_name:
            row += 1
            continue

        b_value = resolve_b_value(wb=wb, ws=ws, row=row, cache=h_cache, stack=set())
        h_value = compute_h_value(wb=wb, ws=ws, row=row, cache=h_cache, stack=set())
        crossed_thresholds = reached_honors(b_value, h_value, thresholds)
        for threshold in crossed_thresholds:
            honored.append((full_name, threshold))

        row += 1

    return honored


def replace_placeholders_in_docx_xml(
    template_file: Path,
    target_file: Path,
    replacements: list[tuple[str, str]],
) -> None:
    def replace_in_text_nodes(xml_text: str) -> str:
        # Nur sichtbaren Text in Word-Textknoten ersetzen, nicht in Attributen/Tags.
        pattern = re.compile(r"(<(?:w|a):t[^>]*>)(.*?)(</(?:w|a):t>)", flags=re.DOTALL)

        def _node_repl(match: re.Match[str]) -> str:
            prefix, content, suffix = match.groups()
            new_content = content
            for old, new in replacements:
                new_content = re.sub(re.escape(old), escape(new), new_content, flags=re.IGNORECASE)
            return f"{prefix}{new_content}{suffix}"

        return pattern.sub(_node_repl, xml_text)

    with zipfile.ZipFile(template_file, "r") as zin, zipfile.ZipFile(target_file, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename.endswith(".xml"):
                text = data.decode("utf-8")
                text = replace_in_text_nodes(text)
                data = text.encode("utf-8")
            zout.writestr(info, data)


def create_certificates(
    template_file: Path,
    output_dir: Path,
    players: list[tuple[str, int]],
    certificate_date: str,
) -> int:
    if not template_file.exists():
        raise FileNotFoundError(f"Vorlage nicht gefunden: {template_file}")

    output_dir.mkdir(parents=True, exist_ok=True)

    created = 0
    for full_name, nnn_value in players:
        first_name, last_name = split_name(full_name)
        first_clean = normalize_name_part(first_name)
        last_clean = normalize_name_part(last_name)

        name_block = f"{last_clean}{first_clean}"
        if not name_block:
            name_block = normalize_name_part(full_name) or "Spieler"
        filename_base = f"{name_block}_{nnn_value}"

        full_name_display = " ".join(part for part in [first_name, last_name] if part)

        replacements = [
            ("Nachname, Vorname", full_name_display),
            ("Nachname Vorname", full_name_display),
            ("Vorname Nachname", full_name_display),
            ("Vorname", first_name),
            ("Nachname", last_name),
            ("nnn", str(nnn_value)),
            ("date", certificate_date),
        ]

        filename_docx = make_safe_filename(filename_base) + ".docx"
        target_docx = output_dir / filename_docx
        try:
            if target_docx.exists():
                target_docx.unlink()
            replace_placeholders_in_docx_xml(template_file, target_docx, replacements)
            created += 1
        except PermissionError:
            print(f"Uebersprungen (Datei geoeffnet): {target_docx}")

    return created


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Erzeugt Urkunden fuer alle Spieler mit Ehrung und ersetzt "
            "Platzhalter in einer Word-Vorlage."
        )
    )
    parser.add_argument(
        "--date",
        default=DEFAULT_CERTIFICATE_DATE,
        help=(
            "Datum fuer Platzhalter dd.mm.yyyy, z. B. 05.12.2026 "
            f"(Standard: {DEFAULT_CERTIFICATE_DATE})"
        ),
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help=f"Blattname im Overall-Excel (Standard: {DEFAULT_SHEET})",
    )
    parser.add_argument(
        "--overall",
        default=OVERALL_FILE,
        help=f"Pfad zur Overall-Datei (Standard: {OVERALL_FILE})",
    )
    parser.add_argument(
        "--template",
        default=DEFAULT_TEMPLATE,
        help=f"Pfad zur Word-Vorlage (Standard: {DEFAULT_TEMPLATE})",
    )
    parser.add_argument(
        "--outdir",
        default=DEFAULT_OUTPUT_DIR,
        help=f"Zielordner fuer erzeugte Dateien (Standard: {DEFAULT_OUTPUT_DIR})",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    overall_file = Path(args.overall)
    template_file = Path(args.template)
    output_dir = Path(args.outdir)

    players = iter_honored_players(overall_file=overall_file, sheet_name=args.sheet)
    if not players:
        print("Keine Ehrungen gefunden. Keine Dateien erzeugt.")
        return

    created = create_certificates(
        template_file=template_file,
        output_dir=output_dir,
        players=players,
        certificate_date=args.date,
    )

    print(f"Ehrungen gefunden: {len(players)}")
    print(f"Urkunden erstellt: {created}")
    print(f"Ausgabeordner: {output_dir.resolve()}")


if __name__ == "__main__":
    main()
